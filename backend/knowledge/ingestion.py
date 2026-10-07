from __future__ import annotations

from pathlib import Path
from typing import Any

import fitz
from docx import Document

from backend.rag.retriever import SemanticRetriever


class DocumentIngestionService:
    """Extract supported documents and index them into RAG."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

    def __init__(
        self,
        retriever: SemanticRetriever | None = None,
    ) -> None:
        self.retriever = retriever or SemanticRetriever()

    def ingest_file(
        self,
        file_path: str | Path,
        source: str | None = None,
    ) -> dict[str, Any]:

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"Document not found: {path}"
            )

        if not path.is_file():
            raise ValueError(
                f"Path is not a file: {path}"
            )

        extension = path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported document type: {extension}"
            )

        text = self.extract_text(path)

        if not text.strip():
            raise ValueError(
                f"No readable text found in: {path.name}"
            )

        document_source = (
            str(source).strip()
            if source
            else path.name
        )

        chunks = self.retriever.add_document(
            source=document_source,
            text=text,
        )

        return {
            "source": document_source,
            "filename": path.name,
            "extension": extension,
            "characters": len(text),
            "chunks_created": chunks,
            "status": "INDEXED",
        }

    def extract_text(
        self,
        file_path: str | Path,
    ) -> str:

        path = Path(file_path)
        extension = path.suffix.lower()

        if extension == ".pdf":
            return self._extract_pdf(path)

        if extension == ".docx":
            return self._extract_docx(path)

        if extension == ".txt":
            return self._extract_txt(path)

        raise ValueError(
            f"Unsupported document type: {extension}"
        )

    def _extract_pdf(
        self,
        path: Path,
    ) -> str:

        parts: list[str] = []

        with fitz.open(path) as document:
            for page_number, page in enumerate(
                document,
                start=1,
            ):
                text = page.get_text("text").strip()

                if text:
                    parts.append(
                        f"[Page {page_number}]\n{text}"
                    )

        return "\n\n".join(parts).strip()

    def _extract_docx(
        self,
        path: Path,
    ) -> str:

        document = Document(path)

        parts: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:

                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    parts.append(
                        " | ".join(cells)
                    )

        return "\n".join(parts).strip()

    def _extract_txt(
        self,
        path: Path,
    ) -> str:

        return path.read_text(
            encoding="utf-8"
        ).strip()

    def search(
        self,
        query: str,
        top_k: int = 3,
        min_score: float = 0.0,
    ) -> list[dict]:

        return self.retriever.retrieve(
            query=query,
            top_k=top_k,
            min_score=min_score,
        )

    def save_index(
        self,
        directory: str | Path,
    ) -> None:

        self.retriever.save(directory)

    def load_index(
        self,
        directory: str | Path,
    ) -> None:

        self.retriever.load(directory)

    def stats(self) -> dict[str, int]:

        return {
            "documents": self.retriever.document_count,
            "chunks": self.retriever.chunk_count,
        }

